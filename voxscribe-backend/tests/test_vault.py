import io
from docx import Document
from tests.conftest import OTHER_USER_ID, USER_ID

SAMPLE = "This is how I write. Short sentences. Lots of energy. " * 20

def _docx_bytes(text: str) -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

def test_list_voices_creates_default(client, fake_db):
    response = client.get("/voices")
    assert response.status_code == 200
    voices = response.json()
    assert [v["name"] for v in voices] == ["Default"]
    assert fake_db.tables["voices"][0]["user_id"] == USER_ID

def test_upload_txt_creates_source_chunks_and_profile(client, fake_db, fake_llm):
    response = client.post("/vault/upload", files={"file": ("notes.txt", SAMPLE.encode())})
    body = response.json()
    assert response.status_code == 200, body
    assert body["status"] == "Success"
    assert body["saved_chunks"] > 0

    source = fake_db.tables["vault_sources"][0]
    assert source["name"] == "notes.txt"
    assert source["source_type"] == "txt"
    assert source["chunk_count"] == body["saved_chunks"]
    assert all(d["source_id"] == source["id"] for d in fake_db.tables["documents"])

    # Background task refreshed the voice's style profile
    assert fake_db.tables["voices"][0]["style_profile"].startswith("- Short punchy")

def test_upload_docx(client, fake_db):
    response = client.post("/vault/upload", files={"file": ("essay.docx", _docx_bytes(SAMPLE))})
    assert response.json()["status"] == "Success"
    assert fake_db.tables["vault_sources"][0]["source_type"] == "docx"

def test_upload_rejects_unsupported_type(client):
    response = client.post("/vault/upload", files={"file": ("image.png", b"\x89PNG")})
    assert response.status_code == 400
    assert "Unsupported" in response.json()["detail"]

def test_add_pasted_text_to_named_voice(client, fake_db):
    voice = client.post("/voices", json={"name": "Company blog"}).json()
    response = client.post("/vault/text", json={"title": "My intro", "text": SAMPLE, "voice_id": voice["id"]})
    assert response.json()["voice_id"] == voice["id"]

    sources = client.get(f"/vault/sources?voice_id={voice['id']}").json()
    assert [s["name"] for s in sources] == ["My intro"]

    listed = {v["name"]: v for v in client.get("/voices").json()}
    assert listed["Company blog"]["source_count"] == 1
    assert listed["Company blog"]["chunk_count"] == sources[0]["chunk_count"]

def test_delete_source_removes_its_chunks(client, fake_db):
    source_id = client.post("/vault/text", json={"text": SAMPLE}).json()["source_id"]
    assert fake_db.tables["documents"]

    assert client.delete(f"/vault/sources/{source_id}").status_code == 204
    assert fake_db.tables["documents"] == []
    assert client.get("/vault/sources").json() == []

def test_delete_voice_cascades(client, fake_db):
    voice = client.post("/voices", json={"name": "Temp"}).json()
    client.post("/vault/text", json={"text": SAMPLE, "voice_id": voice["id"]})

    assert client.delete(f"/voices/{voice['id']}").status_code == 204
    assert fake_db.tables["vault_sources"] == []
    assert fake_db.tables["documents"] == []

def test_cannot_touch_another_users_data(client, fake_db):
    other_voice = fake_db.table("voices").insert({"user_id": OTHER_USER_ID, "name": "Theirs"}).execute().data[0]
    other_source = fake_db.table("vault_sources").insert(
        {"user_id": OTHER_USER_ID, "voice_id": other_voice["id"], "name": "x", "source_type": "pdf", "chunk_count": 1}
    ).execute().data[0]

    assert client.delete(f"/vault/sources/{other_source['id']}").status_code == 404
    assert client.delete(f"/voices/{other_voice['id']}").status_code == 404
    assert client.post("/vault/text", json={"text": SAMPLE, "voice_id": other_voice["id"]}).status_code == 404
    assert len(fake_db.tables["vault_sources"]) == 1

def test_invalid_ids_return_404(client):
    assert client.delete("/vault/sources/not-a-uuid").status_code == 404
    assert client.patch("/voices/not-a-uuid", json={"name": "x"}).status_code == 404

def test_url_import_blocks_private_addresses(client):
    response = client.post("/vault/url", json={"url": "http://127.0.0.1:8000/"})
    assert response.status_code == 400
    assert "private" in response.json()["detail"]

def test_refresh_profile(client, fake_db, fake_llm):
    client.post("/vault/text", json={"text": SAMPLE})
    voice_id = fake_db.tables["voices"][0]["id"]
    fake_llm.response = "- New profile"
    response = client.post(f"/voices/{voice_id}/profile")
    assert response.json() == {"style_profile": "- New profile"}
