import io
import json
from types import SimpleNamespace
import jwt
import pytest
from docx import Document
from app.services import markdown_utils, post_service, publish_service
from tests.conftest import USER_ID

POST = """# Ten Lessons From Remote Work

Remote work changed how we collaborate. It is **not** always easy.

## Communicate clearly

- Write things down
- Default to async
"""

class JsonLLM:
    def __init__(self, payload):
        self.payload = payload
        self.prompts = []

    def invoke(self, messages):
        self.prompts.append(messages[-1].content)
        content = self.payload if isinstance(self.payload, str) else json.dumps(self.payload)
        return SimpleNamespace(content=content)

def test_markdown_utils():
    assert markdown_utils.extract_title(POST) == "Ten Lessons From Remote Work"
    assert not markdown_utils.remove_title(POST).startswith("#")
    assert markdown_utils.slugify("Héllo, World! 2026") == "hello-world-2026"
    assert "<strong>not</strong>" in markdown_utils.to_html(POST)
    stats = markdown_utils.readability(POST)
    assert stats["words"] > 10 and stats["reading_minutes"] == 1 and 0 <= stats["flesch_score"] <= 100

def test_rewrite_uses_voice_profile(client, fake_db, monkeypatch):
    voice = client.post("/voices", json={"name": "Mine"}).json()
    fake_db.tables["voices"][0]["style_profile"] = "- Loves em dashes"
    llm = JsonLLM("Rewritten passage")
    monkeypatch.setattr(post_service, "get_llm", lambda **_: llm)

    response = client.post("/posts/rewrite", json={"text": "Old text", "instruction": "punchier", "voice_id": voice["id"]})
    assert response.json() == {"text": "Rewritten passage"}
    assert "Loves em dashes" in llm.prompts[0] and '"punchier"' in llm.prompts[0]

def test_seo_suggestions_with_internal_links(client, fake_db, monkeypatch):
    earlier = fake_db.table("blog_posts").insert(
        {"user_id": USER_ID, "youtube_url": "x", "content": "# Async Communication Guide\n\nBody"}
    ).execute().data[0]
    current = fake_db.table("blog_posts").insert({"user_id": USER_ID, "youtube_url": "y", "content": POST}).execute().data[0]

    llm = JsonLLM({
        "meta_description": "Learn ten lessons from remote work.",
        "tags": ["#Remote", "Productivity"],
        "internal_links": [{"index": 0, "anchor_text": "Default to async"}, {"index": 99}],
    })
    monkeypatch.setattr(post_service, "get_llm", lambda **_: llm)

    data = client.post("/posts/seo", json={"content": POST, "keyword": "remote work", "post_id": current["id"]}).json()
    assert data["slug"] == "ten-lessons-from-remote-work"
    assert data["tags"] == ["remote", "productivity"]
    assert data["internal_links"] == [{"post_id": earlier["id"], "title": "Async Communication Guide", "anchor_text": "Default to async"}]
    assert "Ten Lessons" not in llm.prompts[0].split("Earlier posts by the same author:")[1].split("Post:")[0]

def test_seo_falls_back_when_llm_returns_garbage(client, fake_db, monkeypatch):
    monkeypatch.setattr(post_service, "get_llm", lambda **_: JsonLLM("no json here"))
    data = client.post("/posts/seo", json={"content": POST}).json()
    assert data["meta_description"].startswith("Remote work changed")
    assert data["tags"] == [] and data["internal_links"] == []

def test_docx_export(client):
    response = client.post("/posts/export/docx", json={"content": POST, "filename": "My Post!"})
    assert response.status_code == 200
    assert 'filename="my-post.docx"' in response.headers["content-disposition"]
    texts = [p.text for p in Document(io.BytesIO(response.content)).paragraphs]
    assert "Ten Lessons From Remote Work" in texts and "Write things down" in texts

@pytest.fixture
def captured(monkeypatch):
    calls = []

    def fake_post(url, json=None, headers=None, timeout=None):
        calls.append({"url": url, "json": json, "headers": headers})
        if "dev.to" in url:
            body = {"id": 1, "url": "https://dev.to/me/post"}
        elif "wp-json" in url:
            body = {"id": 7, "link": "https://blog.example.com/?p=7"}
        elif "ghost" in url:
            body = {"posts": [{"id": "g1", "url": "https://ghost.example.com/p/"}]}
        else:
            body = {"data": {"createDraft": {"draft": {"id": "d1", "slug": "s"}}}}
        return SimpleNamespace(status_code=200, json=lambda: body, text="")

    monkeypatch.setattr(publish_service.httpx, "post", fake_post)
    monkeypatch.setattr(publish_service, "ensure_public_host", lambda url: None)
    return calls

def test_publish_devto_as_draft(client, captured):
    response = client.post("/publish/devto", json={"content": POST, "api_key": "k", "tags": ["remote-work", "a", "b", "c", "d"]})
    assert response.json() == {"platform": "devto", "status": "draft", "url": "https://dev.to/me/post", "id": "1"}
    article = captured[0]["json"]["article"]
    assert article["published"] is False
    assert article["title"] == "Ten Lessons From Remote Work"
    assert not article["body_markdown"].startswith("# Ten")
    assert article["tags"] == ["remotework", "a", "b", "c"]

def test_publish_wordpress_sends_html_and_basic_auth(client, captured):
    client.post("/publish/wordpress", json={"content": POST, "api_key": "app pass", "site_url": "blog.example.com", "username": "me"})
    call = captured[0]
    assert call["url"] == "https://blog.example.com/wp-json/wp/v2/posts"
    assert call["headers"]["Authorization"].startswith("Basic ")
    assert "<h2>Communicate clearly</h2>" in call["json"]["content"]
    assert call["json"]["status"] == "draft"

def test_publish_ghost_signs_jwt(client, captured):
    key = "abc123:" + "11" * 32
    client.post("/publish/ghost", json={"content": POST, "api_key": key, "site_url": "https://ghost.example.com"})
    token = captured[0]["headers"]["Authorization"].split(" ", 1)[1]
    assert jwt.get_unverified_header(token)["kid"] == "abc123"
    jwt.decode(token, bytes.fromhex("11" * 32), algorithms=["HS256"], audience="/admin/")

def test_publish_hashnode_draft(client, captured):
    response = client.post("/publish/hashnode", json={"content": POST, "api_key": "pat", "publication_id": "pub1"})
    assert response.json()["url"] == "https://hashnode.com/draft/d1"
    assert "createDraft" in captured[0]["json"]["query"]

def test_publish_validation_errors(client, captured):
    assert client.post("/publish/medium", json={"content": POST, "api_key": "k"}).status_code == 404
    response = client.post("/publish/wordpress", json={"content": POST, "api_key": "k", "site_url": "blog.example.com"})
    assert response.status_code == 400 and "username" in response.json()["detail"]
    assert client.post("/publish/ghost", json={"content": POST, "api_key": "bad", "site_url": "x.com"}).status_code == 400

def test_publish_blocks_private_sites(client):
    response = client.post(
        "/publish/wordpress",
        json={"content": POST, "api_key": "k", "site_url": "http://127.0.0.1", "username": "me"},
    )
    assert response.status_code == 400 and "private" in response.json()["detail"]
