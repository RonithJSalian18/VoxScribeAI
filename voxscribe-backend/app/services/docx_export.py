"""Converts generated Markdown into a Word document."""
import io
import re
from docx import Document
from docx.shared import Pt, RGBColor

INLINE_PATTERN = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`|\[[^\]]+\]\([^)]+\))")

def _add_inline(paragraph, text: str) -> None:
    """Adds text to a paragraph, applying **bold**, *italic*, `code` and [links](url)."""
    for part in INLINE_PATTERN.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
            paragraph.add_run(f" ({url})")
        elif (part.startswith("*") and part.endswith("*")) or (part.startswith("_") and part.endswith("_")):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)

def markdown_to_docx(content: str) -> bytes:
    document = Document()
    document.styles["Normal"].font.size = Pt(11)

    in_code = False
    code_lines: list[str] = []
    paragraph_lines: list[str] = []

    def flush_paragraph():
        if paragraph_lines:
            _add_inline(document.add_paragraph(), " ".join(l.strip() for l in paragraph_lines))
            paragraph_lines.clear()

    for line in content.splitlines():
        if line.strip().startswith("```"):
            flush_paragraph()
            if in_code:
                para = document.add_paragraph()
                run = para.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                code_lines.clear()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue

        stripped = line.strip()
        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        bullet = re.match(r"^[-*+]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)

        if not stripped:
            flush_paragraph()
        elif heading:
            flush_paragraph()
            level = len(heading.group(1))
            _add_inline(document.add_heading(level=0 if level == 1 else min(level - 1, 9)), heading.group(2))
        elif re.fullmatch(r"([-*_]\s*){3,}", stripped):
            flush_paragraph()
            document.add_paragraph("―" * 20)
        elif bullet:
            flush_paragraph()
            _add_inline(document.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            flush_paragraph()
            _add_inline(document.add_paragraph(style="List Number"), numbered.group(1))
        elif stripped.startswith(">"):
            flush_paragraph()
            _add_inline(document.add_paragraph(style="Intense Quote"), stripped.lstrip("> ").strip())
        else:
            paragraph_lines.append(stripped)

    flush_paragraph()
    if in_code and code_lines:
        document.add_paragraph("\n".join(code_lines))

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
